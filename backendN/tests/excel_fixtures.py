"""Generated XLSX/XLS fixtures; no binary workbooks or private data in Git."""
from datetime import datetime
import io
import zipfile


def make_dense_book(ext, sheets, rows):
    """Benchmark fixture generation is separate from measured conversion."""
    output = io.BytesIO()
    if ext == "xlsx":
        from openpyxl import Workbook
        book = Workbook(write_only=True)
        for index in range(sheets):
            sheet = book.create_sheet(f"Sheet-{index}")
            sheet.append(["id", "name", "amount", "note"])
            for row in range(rows):
                sheet.append([row, f"sheet-{index}-row-{row}", row * 3.5, f"generated café 日本語 {index}-{row}"])
        book.save(output)
        book.close()
    else:
        import xlwt
        book = xlwt.Workbook()
        for index in range(sheets):
            sheet = book.add_sheet(f"Sheet-{index}")
            for column, value in enumerate(["id", "name", "amount", "note"]):
                sheet.write(0, column, value)
            for row in range(rows):
                for column, value in enumerate([row, f"sheet-{index}-row-{row}", row * 3.5, f"generated café 日本語 {index}-{row}"]):
                    sheet.write(row + 1, column, value)
        book.save(output)
    return output.getvalue()


def make_book(ext="xlsx", *, legacy_view=False, image=False):
    sheets = [
        ("Mixed 日本語", [["Name", "Name", "Amount", "Date", "Flag"],
          ["café | pipe", "line\nbreak", 1.23456789, datetime(2024, 2, 3, 4, 5), True],
          ["<tag>&", "", 1e-9, None, False], ["NaN", "NULL", -1000000000.5],
          [None] * 5, ["last-first-sheet", 'quote"value', 123.0]]),
        ("Empty", []),
        ("Results", [["id", "value"], [1, "before blank"], [], [2, "終わり🙂"]]),
    ]
    output = io.BytesIO()
    if ext == "xls":
        import xlwt
        book = xlwt.Workbook()
        date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM:SS")
        for name, rows in sheets:
            sheet = book.add_sheet(name)
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    if value is not None:
                        sheet.write(r, c, value, date_style if isinstance(value, datetime) else xlwt.Style.default_style)
        book.save(output)
        return output.getvalue()
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book["Results"]["D2"] = "=SUM(A2:A4)"
    book["Results"].merge_cells("F2:G2")
    book["Results"]["F2"] = "merged value"
    book["Empty"].sheet_state = "hidden"
    if image:
        from PIL import Image
        from openpyxl.drawing.image import Image as WorkbookImage
        pixels = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(pixels, format="PNG")
        pixels.seek(0)
        book["Results"].add_image(WorkbookImage(pixels), "A8")
    book.save(output)
    book.close()
    if not legacy_view:
        return output.getvalue()
    repaired_fixture = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source, zipfile.ZipFile(repaired_fixture, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = data.replace(b'<sheetView ', b'<sheetView showZeroes="1" ')
            target.writestr(item, data)
    return repaired_fixture.getvalue()
